from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from schema import ExperimentRunRequest, ExperimentRunResponse, Experiment
import random
import time
import numpy as np
from pathlib import Path
import json
import tempfile
import os
from filelock import FileLock
from export_excel import export_results_to_excel
from abc_algorithm import run_abc_experiment

app = FastAPI(title="Bee Algorithm API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000"],
    allow_methods=["*"],
    allow_headers=["*"],
)

DATA_FILE = Path(__file__).parent / 'experiments.json'
LOCK_FILE = str(DATA_FILE) + '.lock'


def read_experiments():
    if not DATA_FILE.exists():
        return []
    with open(DATA_FILE, 'r', encoding='utf-8') as f:
        return json.load(f)


def write_experiments(data):
    lock = FileLock(LOCK_FILE)
    with lock:
        tmp = tempfile.NamedTemporaryFile('w', delete=False, dir=str(DATA_FILE.parent), encoding='utf-8')
        json.dump(data, tmp, ensure_ascii=False, indent=2)
        tmp.flush()
        tmp.close()
        os.replace(tmp.name, DATA_FILE)


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/run", response_model=ExperimentRunResponse)
def run_experiment(req: ExperimentRunRequest):
    """
    Run a real Bee Algorithm experiment using the ABC implementation.
    """
    start_time = time.time()

    # Extract matrix from input
    matrix = req.input.matrix
    if not matrix or len(matrix) == 0:
        raise HTTPException(status_code=400, detail="No matrix data provided")

    # Get matrix dimensions
    n_alternatives = len(matrix)
    n_criteria = len(matrix[0])

    # Run the real ABC algorithm
    try:
        results = run_abc_experiment(
            matrix=matrix,
            num_bees=req.params.numBees,
            iterations=req.params.iterations,
            seed=req.params.seed,
            lb=req.params.lowerBound if req.params.lowerBound is not None else 0.0,
            ub=req.params.upperBound if req.params.upperBound is not None else 1.0
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"ABC algorithm failed: {str(e)}")

    # Extract results
    best_solution = results['bestSolution']
    result_series = results['resultSeries']
    convergence = results['convergence']
    final_fitness = results['bestFitness']

    # Calculate KPIs
    kpis = [
        {"label": "Best fitness", "value": round(final_fitness, 6)},
        {"label": "Iterations", "value": req.params.iterations},
        {"label": "Convergence", "value": round(convergence, 6)},
        {"label": "Alternatives", "value": n_alternatives},
        {"label": "Criteria", "value": n_criteria},
        {"label": "Bees", "value": req.params.numBees},
        {"label": "Trial Limit", "value": req.params.numBees * n_criteria}
    ]

    duration_ms = int((time.time() - start_time) * 1000)

    return ExperimentRunResponse(
        durationMs=duration_ms,
        kpis=kpis,
        bestSolution=best_solution,
        resultSeries=result_series
    )


@app.get("/experiments")
def list_experiments():
    return read_experiments()


@app.get("/experiments/{experiment_id}")
def get_experiment(experiment_id: str):
    """
    Get a single experiment by ID.
    """
    experiments = read_experiments()
    experiment = next((e for e in experiments if e.get('id') == experiment_id), None)
    
    if not experiment:
        raise HTTPException(status_code=404, detail="Experiment not found")
    
    return experiment


@app.post("/experiments", status_code=201)
def create_experiment(exp: Experiment):
    data = read_experiments()
    data.append(exp.dict())
    write_experiments(data)
    return exp.dict()


@app.get("/export/{experiment_id}/excel")
def export_experiment_excel(experiment_id: str):
    """
    Export an experiment's results to Excel format using the template.
    """
    experiments = read_experiments()
    experiment = next((e for e in experiments if e.get('id') == experiment_id), None)
    
    if not experiment:
        raise HTTPException(status_code=404, detail="Experiment not found")
    
    # Transform experiment data to the required Excel format
    rows = []
    result_series = experiment.get('resultSeries', [])
    params = experiment.get('params', {})
    input_data = experiment.get('input', {})
    best_solution = experiment.get('bestSolution', [])
    
    # Calculate trial limit (auto N*D if not specified)
    matrix = input_data.get('matrix', [])
    n_criteria = len(matrix[0]) if matrix and len(matrix) > 0 else 0
    trial_limit = params.get('feedLimit', params.get('numBees', 20) * n_criteria if n_criteria > 0 else 10)
    
    total_duration_sec = experiment.get('durationMs', 0) / 1000.0
    time_per_iteration = total_duration_sec / len(result_series) if result_series else 0
    
    for idx, r in enumerate(result_series):
        iteration = r.get('iteration', idx + 1)
        best_fitness = r.get('bestFitness', 0.0)
        avg_fitness = r.get('avgFitness', 0.0)
        std_fitness = r.get('stdFitness', 0.0)
        worst_fitness = r.get('worstFitness', std_fitness)
        diversity = r.get('diversity', std_fitness)
        num_scouts = r.get('numScouts', 0)
        improvement = r.get('improvement', 1 if idx == 0 else 0)
        
        rows.append({
            'ExperimentName': experiment.get('name', experiment_id),
            'Iteration': iteration,
            'Fbest': best_fitness,
            'Xbest': best_solution if best_solution else [],
            'MeanFitness': avg_fitness,
            'WorstFitness': worst_fitness,
            'NumScouts': num_scouts,
            'Diversity': diversity,
            'Improvement': improvement,
            'Time (s)': round(time_per_iteration, 6),
            'NumBees': params.get('numBees', 20),
            'TrialLimit': trial_limit,
            'Seed': params.get('seed', '')
        })
    
    if not rows:
        raise HTTPException(status_code=400, detail="No data to export")
    
    # Use temporary files for template and output
    temp_dir = Path(tempfile.gettempdir())
    template_path = temp_dir / "tabla_abc_template.xlsx"
    output_path = temp_dir / f"experiment_{experiment_id}_export.xlsx"
    
    # Create a minimal template if it doesn't exist
    if not template_path.exists():
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Datos"
        
        # Write headers with formatting
        headers = [
            "ExperimentName", "Iteration", "Fbest", "Xbest", "MeanFitness", 
            "WorstFitness", "NumScouts", "Diversity", "Improvement", 
            "Time (s)", "NumBees", "TrialLimit", "Seed"
        ]
        
        # Header row styling
        header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        header_font = Font(bold=True)
        
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Set column widths for readability
        ws.column_dimensions['A'].width = 20  # ExperimentName
        ws.column_dimensions['B'].width = 10  # Iteration
        ws.column_dimensions['C'].width = 12  # Fbest
        ws.column_dimensions['D'].width = 30  # Xbest
        ws.column_dimensions['E'].width = 12  # MeanFitness
        ws.column_dimensions['F'].width = 12  # WorstFitness
        ws.column_dimensions['G'].width = 10  # NumScouts
        ws.column_dimensions['H'].width = 10  # Diversity
        ws.column_dimensions['I'].width = 12  # Improvement
        ws.column_dimensions['J'].width = 10  # Time (s)
        ws.column_dimensions['K'].width = 10  # NumBees
        ws.column_dimensions['L'].width = 10  # TrialLimit
        ws.column_dimensions['M'].width = 10  # Seed
        
        wb.save(str(template_path))
    
    try:
        export_results_to_excel(
            rows,
            template_path=str(template_path),
            output_path=str(output_path)
        )
        
        return FileResponse(
            path=str(output_path),
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            filename=f"experiment_{experiment_id}.xlsx"
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Excel export failed: {str(e)}")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8001)
